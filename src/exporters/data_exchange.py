"""Portable questionnaire import/export formats with explicit rights-neutral defaults."""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from schemas.questionnaire_schema import (
    ItemSchema,
    QuestionnaireMetadata,
    QuestionnaireParent,
    QuestionnaireVersion,
    ResponseOption,
)
from src.exporters.r_syntax_gen import export_r_syntax

FHIR_QUESTIONNAIRE_R4 = "http://hl7.org/fhir/StructureDefinition/Questionnaire"
FHIR_DIMENSION_EXTENSION = "https://psymetriq.org/fhir/StructureDefinition/item-dimension"
FHIR_VARIABLE_NAME_EXTENSION = "https://psymetriq.org/fhir/StructureDefinition/variable-name"
FHIR_RESPONSE_MODE_EXTENSION = "https://psymetriq.org/fhir/StructureDefinition/response-mode"
FHIR_RESPONSE_SET_EXTENSION = "https://psymetriq.org/fhir/StructureDefinition/response-set"
FHIR_VERSION_DISPLAY_EXTENSION = "https://psymetriq.org/fhir/StructureDefinition/version-display-name"
FHIR_NUMERIC_RESPONSE_CODE_EXTENSION = (
    "https://psymetriq.org/fhir/StructureDefinition/numeric-response-code"
)
FHIR_ITEM_METADATA_NOTE_EXTENSION = (
    "https://psymetriq.org/fhir/StructureDefinition/item-metadata-note"
)
XLSX_SHEETS = ("Items", "Antwortoptionen", "Scoring", "Quellen")

REDCAP_HEADERS = [
    "Variable / Field Name",
    "Form Name",
    "Section Header",
    "Field Type",
    "Field Label",
    "Choices, Calculations, OR Slider Labels",
    "Field Note",
    "Text Validation Type OR Show Slider Number",
    "Text Validation Min",
    "Text Validation Max",
    "Identifier?",
    "Branching Logic (Show field only if...)",
    "Required Field?",
    "Custom Alignment",
    "Question Number (surveys only)",
    "Matrix Group Name",
    "Matrix Ranking?",
    "Field Annotation",
]

# One source of truth mapping REDCap's Data Dictionary CSV header row to the
# snake_case keys the REDCap API's metadata import/export JSON uses for the
# same columns, so the CSV writer/reader and the live-API path (see
# src/exporters/redcap_api.py) can never silently drift apart.
REDCAP_HEADER_TO_KEY: dict[str, str] = {
    "Variable / Field Name": "field_name",
    "Form Name": "form_name",
    "Section Header": "section_header",
    "Field Type": "field_type",
    "Field Label": "field_label",
    "Choices, Calculations, OR Slider Labels": "select_choices_or_calculations",
    "Field Note": "field_note",
    "Text Validation Type OR Show Slider Number": "text_validation_type_or_show_slider_number",
    "Text Validation Min": "text_validation_min",
    "Text Validation Max": "text_validation_max",
    "Identifier?": "identifier",
    "Branching Logic (Show field only if...)": "branching_logic",
    "Required Field?": "required_field",
    "Custom Alignment": "custom_alignment",
    "Question Number (surveys only)": "question_number",
    "Matrix Group Name": "matrix_group_name",
    "Matrix Ranking?": "matrix_ranking",
    "Field Annotation": "field_annotation",
}
REDCAP_KEY_TO_HEADER: dict[str, str] = {key: header for header, key in REDCAP_HEADER_TO_KEY.items()}
REDCAP_RESPONDENT_FIELD_TYPES = frozenset(
    {"radio", "dropdown", "checkbox", "text", "notes", "yesno", "truefalse", "slider"}
)
REDCAP_NON_RESPONDENT_FIELD_TYPES = frozenset({"calc", "descriptive", "file", "sql"})


class DataExchangeError(ValueError):
    """Raised when external questionnaire data is malformed or unsupported."""


def _safe_identifier(value: str, fallback: str = "instrument") -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip()).strip("_").lower()
    if not normalized or not normalized[0].isalpha():
        normalized = f"{fallback}_{normalized}" if normalized else fallback
    return normalized[:64].rstrip("_") or fallback


def _safe_variable_name(value: str, index: int) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "_", value.strip()).strip("_")
    if not normalized or not normalized[0].isalpha():
        normalized = f"item_{normalized}" if normalized else f"item_{index:03d}"
    return normalized[:26].rstrip("_") or f"item_{index:03d}"


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def import_psymetriq_json(content: str) -> list[QuestionnaireParent]:
    """Validate one or multiple canonical instrument families from JSON text."""
    try:
        data = json.loads(content)
        if isinstance(data, dict) and "questionnaires" in data:
            data = data["questionnaires"]
        candidates = data if isinstance(data, list) else [data]
        if not candidates:
            raise DataExchangeError("The JSON file contains no questionnaire records")
        questionnaires = [QuestionnaireParent.model_validate(item) for item in candidates]
    except (json.JSONDecodeError, TypeError, ValidationError) as error:
        raise DataExchangeError(
            f"The file is not valid PsyMetriQ questionnaire JSON ({type(error).__name__})"
        ) from None
    identifiers = [questionnaire.instrument_id for questionnaire in questionnaires]
    if len(identifiers) != len(set(identifiers)):
        raise DataExchangeError("The import contains duplicate instrument_id values")
    return questionnaires


def export_psymetriq_json(questionnaires: list[QuestionnaireParent]) -> str:
    """Serialize validated instrument families in the canonical portable schema."""
    return json.dumps(
        {
            "format": "psymetriq-questionnaire-catalog",
            "schema_version": 1,
            "questionnaires": [
                questionnaire.model_dump(mode="json") for questionnaire in questionnaires
            ],
        },
        ensure_ascii=False,
        indent=2,
    ) + "\n"


def _fhir_answer_option(option: ResponseOption) -> dict[str, Any]:
    coding: dict[str, Any] = {"code": str(option.code), "display": option.label}
    value: dict[str, Any] = {"valueCoding": coding}
    if isinstance(option.code, int):
        coding["extension"] = [
            {"url": FHIR_NUMERIC_RESPONSE_CODE_EXTENSION, "valueInteger": option.code}
        ]
    if option.score is not None:
        coding.setdefault("extension", []).append(
            {
                "url": "https://psymetriq.org/fhir/StructureDefinition/response-score",
                "valueDecimal": option.score,
            }
        )
    return value


def export_fhir_questionnaire(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion
) -> dict[str, Any]:
    """Export one concrete instrument version as an HL7 FHIR R4 Questionnaire."""
    items: list[dict[str, Any]] = []
    for item in version.items:
        fhir_type = {
            "categorical": "choice",
            "numeric": "decimal",
            "text": "string",
        }[item.response_mode]
        fhir_item: dict[str, Any] = {
            "linkId": item.item_id,
            "text": item.prompt_text,
            "type": fhir_type,
            "required": item.is_required,
            "extension": [
                {"url": FHIR_DIMENSION_EXTENSION, "valueString": item.dimension},
                {"url": FHIR_VARIABLE_NAME_EXTENSION, "valueString": item.variable_name},
                {"url": FHIR_RESPONSE_MODE_EXTENSION, "valueCode": item.response_mode},
            ],
        }
        if item.response_set_ref:
            fhir_item["extension"].append(
                {"url": FHIR_RESPONSE_SET_EXTENSION, "valueString": item.response_set_ref}
            )
        if item.response_mode == "categorical" and item.response_set_ref:
            fhir_item["answerOption"] = [
                _fhir_answer_option(option)
                for option in version.response_sets[item.response_set_ref]
            ]
        if item.measurement_unit:
            fhir_item["extension"].append(
                {"url": "https://psymetriq.org/fhir/StructureDefinition/measurement-unit",
                 "valueString": item.measurement_unit}
            )
        if item.metadata.notes:
            fhir_item["extension"].append(
                {
                    "url": FHIR_ITEM_METADATA_NOTE_EXTENSION,
                    "valueString": item.metadata.notes,
                }
            )
        items.append(fhir_item)

    return {
        "resourceType": "Questionnaire",
        "id": _safe_identifier(f"{questionnaire.instrument_id}-{version.version_id}"),
        "meta": {"profile": [FHIR_QUESTIONNAIRE_R4]},
        "identifier": [
            {"system": "https://psymetriq.org/identifier/instrument",
             "value": questionnaire.instrument_id}
        ],
        "url": f"https://psymetriq.org/questionnaires/{questionnaire.instrument_id}/{version.version_id}",
        "version": version.version_id,
        "name": _safe_identifier(questionnaire.instrument_id, "questionnaire"),
        "title": questionnaire.name_full,
        "status": "active",
        "language": version.language,
        "extension": (
            [
                {
                    "url": FHIR_VERSION_DISPLAY_EXTENSION,
                    "valueString": version.display_name,
                }
            ]
            if version.display_name
            else []
        ),
        "item": items,
    }


def import_fhir_questionnaire(resource: dict[str, Any]) -> QuestionnaireParent:
    """Import one FHIR R4 Questionnaire as a rights-unassessed family."""
    if resource.get("resourceType") != "Questionnaire":
        raise DataExchangeError("Expected a FHIR Questionnaire resource")
    flat_items: list[tuple[str, dict[str, Any]]] = []

    def collect(items: Any, section: str = "general") -> None:
        if not isinstance(items, list):
            return
        for entry in items:
            if not isinstance(entry, dict):
                continue
            if entry.get("type") == "group":
                collect(entry.get("item"), entry.get("text") or section)
            elif entry.get("text"):
                flat_items.append((section, entry))

    collect(resource.get("item"))
    if not flat_items:
        raise DataExchangeError("FHIR Questionnaire contains no importable question items")

    version_id = str(resource.get("version") or resource.get("id") or "v1")
    identifier = next(
        (
            entry.get("value")
            for entry in resource.get("identifier", [])
            if isinstance(entry, dict)
            and entry.get("system") == "https://psymetriq.org/identifier/instrument"
            and entry.get("value")
        ),
        None,
    )
    instrument_id = _safe_identifier(
        str(identifier or resource.get("id") or resource.get("title") or "questionnaire")
    )
    version_display_name = next(
        (
            extension.get("valueString")
            for extension in resource.get("extension", [])
            if isinstance(extension, dict)
            and extension.get("url") == FHIR_VERSION_DISPLAY_EXTENSION
            and extension.get("valueString")
        ),
        None,
    )
    response_sets: dict[str, list[ResponseOption]] = {}
    items: list[ItemSchema] = []
    for index, (section, entry) in enumerate(flat_items, start=1):
        extensions = {
            extension.get("url"): extension
            for extension in entry.get("extension", [])
            if isinstance(extension, dict) and extension.get("url")
        }
        response_mode = extensions.get(FHIR_RESPONSE_MODE_EXTENSION, {}).get("valueCode")
        type_name = entry.get("type")
        if response_mode not in ("categorical", "numeric", "text"):
            response_mode = "categorical" if type_name == "choice" else (
                "numeric" if type_name in ("integer", "decimal", "quantity") else "text"
            )
        response_ref: str | None = None
        if response_mode == "categorical":
            response_ref = _safe_identifier(
                extensions.get(FHIR_RESPONSE_SET_EXTENSION, {}).get("valueString")
                or f"scale_{index:03d}",
                "scale",
            )
            if response_ref in response_sets:
                response_ref = f"{response_ref}_{index}"
            options: list[ResponseOption] = []
            for option_index, answer in enumerate(entry.get("answerOption", [])):
                coding = answer.get("valueCoding", {}) if isinstance(answer, dict) else {}
                if not coding:
                    continue
                score_extension = next(
                    (ext for ext in coding.get("extension", [])
                     if ext.get("url", "").endswith("/response-score")),
                    {},
                )
                numeric_code_extension = next(
                    (ext for ext in coding.get("extension", [])
                     if ext.get("url") == FHIR_NUMERIC_RESPONSE_CODE_EXTENSION),
                    {},
                )
                options.append(
                    ResponseOption(
                        code=numeric_code_extension.get(
                            "valueInteger", coding.get("code", option_index)
                        ),
                        label=str(coding.get("display") or coding.get("code") or option_index),
                        score=_number(score_extension.get("valueDecimal")),
                    )
                )
            if not options:
                raise DataExchangeError(
                    f"FHIR choice item {entry.get('linkId', index)!r} has no answer options"
                )
            response_sets[response_ref] = options

        variable_name = _safe_variable_name(
            str(
                extensions.get(FHIR_VARIABLE_NAME_EXTENSION, {}).get("valueString")
                or entry.get("linkId")
                or f"item_{index:03d}"
            ),
            index,
        )
        dimension = extensions.get(FHIR_DIMENSION_EXTENSION, {}).get("valueString") or section
        unit = next(
            (ext.get("valueString") for ext in entry.get("extension", [])
             if isinstance(ext, dict)
             and ext.get("url", "").endswith("/measurement-unit")),
            None,
        )
        item_note = next(
            (
                ext.get("valueString")
                for ext in entry.get("extension", [])
                if isinstance(ext, dict)
                and ext.get("url") == FHIR_ITEM_METADATA_NOTE_EXTENSION
                and ext.get("valueString")
            ),
            None,
        )
        items.append(
            ItemSchema(
                item_id=str(entry.get("linkId") or f"item_{index:03d}"),
                variable_name=variable_name,
                dimension=str(dimension),
                prompt_text=str(entry["text"]),
                response_mode=response_mode,
                response_set_ref=response_ref,
                is_required=bool(entry.get("required", False)),
                measurement_unit=unit,
                metadata=QuestionnaireMetadata(notes=item_note),
                redcap_field_type="checkbox" if type_name == "open-choice" else (
                    "radio" if response_mode == "categorical" else "text"
                ),
            )
        )

    try:
        version = QuestionnaireVersion(
            version_id=version_id,
            language=str(resource.get("language") or "en"),
            display_name=version_display_name,
            response_sets=response_sets,
            items=items,
        )
        return QuestionnaireParent(
            instrument_id=instrument_id,
            name_full=str(resource.get("title") or resource.get("name") or instrument_id),
            is_commercial=None,
            versions=[version],
        )
    except ValidationError as error:
        raise DataExchangeError(
            "FHIR Questionnaire did not meet PsyMetriQ schema requirements "
            f"({type(error).__name__})"
        ) from None


def export_item_csv(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion
) -> str:
    """Export item-level data in a UTF-8 CSV suitable for review and mapping."""
    output = io.StringIO(newline="")
    fields = [
        "instrument_id", "instrument_name", "version_id", "language", "item_id",
        "variable_name", "dimension", "prompt_text", "response_mode", "response_set_ref",
        "response_options", "measurement_unit", "numeric_minimum", "numeric_maximum",
        "is_required", "is_scored", "is_reverse_scored", "metadata_notes",
    ]
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for item in version.items:
        options = (
            version.response_sets[item.response_set_ref]
            if item.response_set_ref is not None
            else []
        )
        writer.writerow(
            {
                "instrument_id": questionnaire.instrument_id,
                "instrument_name": questionnaire.name_full,
                "version_id": version.version_id,
                "language": version.language,
                "item_id": item.item_id,
                "variable_name": item.variable_name,
                "dimension": item.dimension,
                "prompt_text": item.prompt_text,
                "response_mode": item.response_mode,
                "response_set_ref": item.response_set_ref or "",
                "response_options": json.dumps(
                    [option.model_dump(mode="json") for option in options], ensure_ascii=False
                ),
                "measurement_unit": item.measurement_unit or "",
                "numeric_minimum": item.numeric_minimum,
                "numeric_maximum": item.numeric_maximum,
                "is_required": item.is_required,
                "is_scored": item.is_scored,
                "is_reverse_scored": item.is_reverse_scored,
                "metadata_notes": item.metadata.notes or "",
            }
        )
    return output.getvalue()


def _format_redcap_number(value: float) -> str:
    """Render a numeric bound without a spurious trailing ``.0`` for whole numbers."""
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def _format_redcap_code_literal(code: str | int) -> str:
    """Render a response code as a REDCap calc-syntax literal (numeric or quoted string)."""
    if isinstance(code, bool):
        return "1" if code else "0"
    if isinstance(code, int):
        return str(code)
    try:
        numeric = float(code)
    except (TypeError, ValueError):
        return "'" + str(code).replace("'", "\\'") + "'"
    return _format_redcap_number(numeric)


def _redcap_score_expression(
    item: ItemSchema, response_sets: dict[str, list[ResponseOption]]
) -> str:
    """Return the REDCap calc-syntax expression for one item's scored contribution.

    When every response option's stored code already equals its score (the
    common case), the raw field value is used directly; otherwise the option
    scores are recoded explicitly with nested ``if()`` calls, so this always
    reflects ``ResponseOption.score`` rather than REDCap's raw stored choice
    code. ``item.is_reverse_scored`` is applied afterwards, on top of either
    path, using the standard ``(min + max) - value`` reversal over that
    response set's own score range: catalog response sets record each
    option's forward-direction score once and are shared by both
    forward-keyed and reverse-keyed items (see the Rosenberg Self-Esteem
    Scale), so the score field alone does not already encode reversal.
    """
    variable = item.variable_name
    if item.response_mode != "categorical" or not item.response_set_ref:
        return f"[{variable}]"
    options = response_sets.get(item.response_set_ref, [])
    scored_options = [option for option in options if option.score is not None]
    if not scored_options:
        return f"[{variable}]"

    def _code_matches_score(option: ResponseOption) -> bool:
        try:
            return float(option.code) == option.score
        except (TypeError, ValueError):
            return False

    if len(scored_options) == len(options) and all(
        _code_matches_score(option) for option in scored_options
    ):
        expression = f"[{variable}]"
    else:
        expression = "0"
        for option in reversed(scored_options):
            code_literal = _format_redcap_code_literal(option.code)
            score_literal = _format_redcap_number(option.score)
            expression = f"if([{variable}]={code_literal},{score_literal},{expression})"

    if item.is_reverse_scored:
        score_values = [option.score for option in scored_options]
        reversal_sum = _format_redcap_number(min(score_values) + max(score_values))
        expression = f"({reversal_sum}-({expression}))"
    return expression


def _build_redcap_item_records(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion, *, form_name: str
) -> list[dict[str, str]]:
    """Build one REDCap metadata record per item, in the REDCap API's JSON key shape."""
    records: list[dict[str, str]] = []
    for item in version.items:
        field_type = item.redcap_field_type
        validation_type = ""
        minimum = ""
        maximum = ""
        choices = ""
        if item.response_mode == "categorical" and item.response_set_ref:
            field_type = "checkbox" if field_type == "checkbox" else "radio"
            choices = " | ".join(
                f"{option.code}, {option.label}"
                for option in version.response_sets[item.response_set_ref]
            )
        elif item.response_mode == "numeric":
            minimum = (
                "" if item.numeric_minimum is None else _format_redcap_number(item.numeric_minimum)
            )
            maximum = (
                "" if item.numeric_maximum is None else _format_redcap_number(item.numeric_maximum)
            )
            if field_type == "slider":
                slider_minimum = item.numeric_minimum if item.numeric_minimum is not None else 0
                slider_maximum = item.numeric_maximum if item.numeric_maximum is not None else 100
                choices = (
                    f"{_format_redcap_number(slider_minimum)}, , "
                    f"{_format_redcap_number(slider_maximum)}"
                )
            else:
                field_type = "text"
                validation_type = "number"
        else:
            field_type = "notes" if field_type == "slider" else "text"
        records.append(
            {
                "field_name": item.variable_name,
                "form_name": form_name,
                "section_header": item.dimension,
                "field_type": field_type,
                "field_label": item.prompt_text,
                "select_choices_or_calculations": choices,
                "field_note": item.measurement_unit or "",
                "text_validation_type_or_show_slider_number": validation_type,
                "text_validation_min": minimum,
                "text_validation_max": maximum,
                "identifier": "",
                "branching_logic": "",
                "required_field": "y" if item.is_required else "",
                "custom_alignment": "",
                "question_number": "",
                "matrix_group_name": "",
                "matrix_ranking": "",
                "field_annotation": (
                    f"item_id={item.item_id}; scored={str(item.is_scored).lower()}"
                    + (f"; {item.metadata.notes}" if item.metadata.notes else "")
                ),
            }
        )
    return records


def _build_redcap_calc_records(
    version: QuestionnaireVersion, *, form_name: str, used_field_names: set[str]
) -> list[dict[str, str]]:
    """Build one REDCap ``calc`` field per scoring algorithm, deriving true scores.

    ``used_field_names`` (case-insensitive) is extended in place so a calc
    field's name never collides with an item's ``variable_name`` or an
    earlier calc field from the same version.
    """
    records: list[dict[str, str]] = []
    items_by_id = {item.item_id: item for item in version.items}
    for index, algorithm in enumerate(version.scoring_algorithms, start=1):
        ordered_items = [
            items_by_id[item_id] for item_id in algorithm.target_items if item_id in items_by_id
        ]
        if not ordered_items:
            continue
        expressions = [
            _redcap_score_expression(item, version.response_sets) for item in ordered_items
        ]
        total = "(" + "+".join(expressions) + ")"
        if algorithm.method == "mean":
            total = f"({total}/{len(expressions)})"
        if algorithm.multiplier != 1:
            total = f"({total}*{_format_redcap_number(algorithm.multiplier)})"

        base_name = _safe_variable_name(algorithm.output_variable, 900 + index)
        field_name = base_name
        suffix = 1
        while field_name.casefold() in used_field_names:
            suffix += 1
            field_name = _safe_variable_name(f"{base_name}_{suffix}", 900 + index)
        used_field_names.add(field_name.casefold())

        records.append(
            {
                "field_name": field_name,
                "form_name": form_name,
                "section_header": "Berechnete Scores",
                "field_type": "calc",
                "field_label": (
                    f"{algorithm.output_variable} "
                    f"(berechnet: {algorithm.method}, Faktor {algorithm.multiplier:g})"
                ),
                "select_choices_or_calculations": total,
                "field_note": algorithm.missing_data_rules or "",
                "text_validation_type_or_show_slider_number": "",
                "text_validation_min": "",
                "text_validation_max": "",
                "identifier": "",
                "branching_logic": "",
                "required_field": "",
                "custom_alignment": "",
                "question_number": "",
                "matrix_group_name": "",
                "matrix_ranking": "",
                "field_annotation": (
                    f"scoring_algorithm={algorithm.output_variable}; "
                    f"target_items={','.join(algorithm.target_items)}"
                ),
            }
        )
    return records


def build_redcap_metadata_records(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion
) -> list[dict[str, str]]:
    """Build the full REDCap metadata (item + calc-score fields) for one version.

    This is the single source of truth for REDCap export: both the Data
    Dictionary CSV (``export_redcap_data_dictionary``) and the live-API push
    (``src/exporters/redcap_api.py``) are built from these same records, so
    the file you can preview/export locally is exactly what would be pushed.
    """
    form_name = _safe_identifier(questionnaire.instrument_id, "instrument")[:26]
    records = _build_redcap_item_records(questionnaire, version, form_name=form_name)
    used_field_names = {str(record["field_name"]).casefold() for record in records}
    records.extend(
        _build_redcap_calc_records(version, form_name=form_name, used_field_names=used_field_names)
    )
    return records


def export_redcap_data_dictionary(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion
) -> str:
    """Export one version to REDCap's Data Dictionary CSV column convention."""
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=REDCAP_HEADERS, lineterminator="\n")
    writer.writeheader()
    for record in build_redcap_metadata_records(questionnaire, version):
        writer.writerow(
            {header: record.get(key, "") for header, key in REDCAP_HEADER_TO_KEY.items()}
        )
    return output.getvalue()


_REDCAP_YESNO_OPTIONS = [
    ResponseOption(code=1, label="Yes", score=None),
    ResponseOption(code=0, label="No", score=None),
]
_REDCAP_TRUEFALSE_OPTIONS = [
    ResponseOption(code=1, label="True", score=None),
    ResponseOption(code=0, label="False", score=None),
]


def import_redcap_data_dictionary(content: str, language: str) -> QuestionnaireParent:
    """Import a REDCap Data Dictionary CSV without inferring reuse permissions.

    ``radio``/``dropdown``/``checkbox``/``text``/``notes``/``yesno``/
    ``truefalse``/``slider`` fields become respondent items. ``calc``,
    ``descriptive``, ``file``, and ``sql`` fields (and any other type this
    importer does not recognize) are never silently dropped: each one is
    listed by field name and type in the imported family's metadata notes,
    so nothing disappears from the dictionary without a visible trace.
    """
    try:
        reader = csv.DictReader(io.StringIO(content.lstrip("\ufeff")))
        rows = list(reader)
    except csv.Error as error:
        raise DataExchangeError("Could not read REDCap Data Dictionary CSV") from error
    required_headers = {"Variable / Field Name", "Form Name", "Field Type", "Field Label"}
    if not reader.fieldnames or not required_headers.issubset(reader.fieldnames):
        raise DataExchangeError("CSV is missing required REDCap Data Dictionary columns")

    rows = [row for row in rows if (row.get("Field Label") or "").strip()]
    if not rows:
        raise DataExchangeError("REDCap Data Dictionary contains no importable fields")

    skipped_fields: list[str] = []
    importable_rows: list[dict[str, str]] = []
    for row in rows:
        field_type = (row.get("Field Type") or "text").casefold()
        field_name = (row.get("Variable / Field Name") or "").strip() or "(unnamed)"
        if field_type in REDCAP_RESPONDENT_FIELD_TYPES:
            importable_rows.append(row)
        else:
            skipped_fields.append(f"{field_name} ({field_type or 'unknown type'})")
    if not importable_rows:
        raise DataExchangeError("REDCap Data Dictionary contains no importable fields")

    response_sets: dict[str, list[ResponseOption]] = {}
    items: list[ItemSchema] = []
    used_variable_names: set[str] = set()
    for index, row in enumerate(importable_rows, start=1):
        field_type = (row.get("Field Type") or "text").casefold()
        source_name = (row.get("Variable / Field Name") or "").strip()
        variable_name = _safe_variable_name(source_name, index)
        if variable_name.casefold() in used_variable_names:
            variable_name = _safe_variable_name(f"{variable_name}_{index}", index)
        used_variable_names.add(variable_name.casefold())

        validation_type = (
            row.get("Text Validation Type OR Show Slider Number") or ""
        ).casefold()
        is_numeric = field_type == "text" and (
            validation_type in {"number", "integer", "float"}
            or validation_type.startswith("number_")
        )
        mode: Literal["categorical", "numeric", "text"] = (
            "categorical" if field_type in {"radio", "dropdown", "checkbox", "yesno", "truefalse"}
            else "numeric" if is_numeric or field_type == "slider"
            else "text"
        )
        response_ref: str | None = None
        if field_type == "yesno":
            response_ref = f"redcap_scale_{index:03d}"
            response_sets[response_ref] = _REDCAP_YESNO_OPTIONS
        elif field_type == "truefalse":
            response_ref = f"redcap_scale_{index:03d}"
            response_sets[response_ref] = _REDCAP_TRUEFALSE_OPTIONS
        elif mode == "categorical":
            response_ref = f"redcap_scale_{index:03d}"
            options: list[ResponseOption] = []
            for raw_option in (
                row.get("Choices, Calculations, OR Slider Labels") or ""
            ).split("|"):
                code, separator, label = raw_option.strip().partition(",")
                if not separator:
                    continue
                normalized_code: str | int = (
                    int(code.strip())
                    if code.strip().lstrip("-").isdigit()
                    else code.strip()
                )
                options.append(
                    ResponseOption(
                        code=normalized_code,
                        label=label.strip() or str(normalized_code),
                        score=None,
                    )
                )
            if not options:
                raise DataExchangeError(
                    f"REDCap field {source_name or index!r} has no parseable choices"
                )
            response_sets[response_ref] = options

        minimum = _number(row.get("Text Validation Min")) if mode == "numeric" else None
        maximum = _number(row.get("Text Validation Max")) if mode == "numeric" else None
        if field_type == "slider" and minimum is None and maximum is None:
            minimum, maximum = 0.0, 100.0
        item_id = _safe_identifier(source_name or f"item_{index:03d}", "item")
        items.append(
            ItemSchema(
                item_id=item_id,
                variable_name=variable_name,
                dimension=(row.get("Section Header") or row.get("Form Name") or "general").strip(),
                prompt_text=(row.get("Field Label") or "").strip(),
                response_mode=mode,
                response_set_ref=response_ref,
                is_required=(row.get("Required Field?") or "").strip().casefold()
                in {"y", "yes", "1", "true"},
                measurement_unit=(row.get("Field Note") or None),
                numeric_minimum=minimum,
                numeric_maximum=maximum,
                redcap_field_type=(
                    "slider" if field_type == "slider"
                    else field_type if field_type in {"radio", "checkbox"}
                    else "text"
                ),
            )
        )

    form_names = sorted(
        {(row.get("Form Name") or "").strip() for row in importable_rows if row.get("Form Name")}
    )
    instrument_name = (
        form_names[0].replace("_", " ").title()
        if len(form_names) == 1
        else "Imported REDCap Instrument"
    )
    source_id = Path(form_names[0] if len(form_names) == 1 else "redcap_import").stem
    notes = (
        (
            "Fields present in the REDCap Data Dictionary but not imported as "
            "respondent items (unsupported field type, e.g. calc/descriptive/"
            f"file/sql): {'; '.join(skipped_fields)}"
        )
        if skipped_fields
        else None
    )
    return QuestionnaireParent(
        instrument_id=_safe_identifier(source_id),
        name_full=instrument_name,
        is_commercial=None,
        metadata=QuestionnaireMetadata(notes=notes),
        versions=[
            QuestionnaireVersion(
                version_id="redcap_import_v1",
                language=language,
                display_name="Imported from REDCap Data Dictionary",
                response_sets=response_sets,
                items=items,
            )
        ],
    )


def import_questionnaire_file(
    path: Path | str, *, redcap_language: str = "en"
) -> list[QuestionnaireParent]:
    """Load PsyMetriQ JSON, FHIR Questionnaire JSON, or REDCap Data Dictionary CSV."""
    source_path = Path(path)
    try:
        content = source_path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as error:
        raise DataExchangeError(f"Could not read import file {source_path.name}") from error
    return import_questionnaire_content(
        content,
        file_name=source_path.name,
        redcap_language=redcap_language,
    )


def import_questionnaire_content(
    content: str,
    *,
    file_name: str,
    redcap_language: str = "en",
) -> list[QuestionnaireParent]:
    """Import decoded uploaded content when a local filesystem path is unavailable."""
    if Path(file_name).suffix.casefold() == ".csv":
        return [import_redcap_data_dictionary(content, redcap_language)]
    if Path(file_name).suffix.casefold() != ".json":
        raise DataExchangeError("Supported import formats are JSON and REDCap CSV")
    try:
        payload = json.loads(content)
    except json.JSONDecodeError:
        raise DataExchangeError("The JSON file is malformed") from None
    if isinstance(payload, dict) and payload.get("resourceType") == "Questionnaire":
        return [import_fhir_questionnaire(payload)]
    return import_psymetriq_json(content)


def export_questionnaire_xlsx(
    questionnaire: QuestionnaireParent, version: QuestionnaireVersion
) -> bytes:
    """Create a workbook with separate item, response, scoring, and provenance sheets."""
    workbook = Workbook()
    items_sheet = workbook.active
    items_sheet.title = XLSX_SHEETS[0]
    items_sheet.append(
        [
            "Instrument ID", "Instrument", "Version ID", "Version", "Sprache", "Locale",
            "Item ID", "Variable", "Dimension", "Itemtext", "Antwortmodus", "Antwortset",
            "Antwort erforderlich", "Einheit", "Minimum", "Maximum", "Umgekehrt kodiert",
            "Scored", "Metadatenhinweise",
        ]
    )
    for item in version.items:
        items_sheet.append(
            [
                questionnaire.instrument_id, questionnaire.name_full, version.version_id,
                version.display_name or "", version.language, version.locale or "",
                item.item_id, item.variable_name, item.dimension, item.prompt_text,
                item.response_mode, item.response_set_ref or "", item.is_required,
                item.measurement_unit or "", item.numeric_minimum, item.numeric_maximum,
                item.is_reverse_scored, item.is_scored, item.metadata.notes or "",
            ]
        )

    responses_sheet = workbook.create_sheet(XLSX_SHEETS[1])
    responses_sheet.append(["Antwortset", "Code", "Label", "Score"])
    for set_name, options in version.response_sets.items():
        for option in options:
            responses_sheet.append([set_name, option.code, option.label, option.score])

    scores_sheet = workbook.create_sheet(XLSX_SHEETS[2])
    scores_sheet.append(["Score", "Methode", "Item IDs", "Multiplikator", "Missing-Data-Regel"])
    for algorithm in version.scoring_algorithms:
        scores_sheet.append(
            [
                algorithm.output_variable, algorithm.method, ", ".join(algorithm.target_items),
                algorithm.multiplier, algorithm.missing_data_rules or "",
            ]
        )

    source_sheet = workbook.create_sheet(XLSX_SHEETS[3])
    source_sheet.append(
        ["Titel", "Dokumenttyp", "Sprache", "Quelle", "Lizenz", "Weitergabe erlaubt", "Basis"]
    )
    for source in version.source_documents:
        source_sheet.append(
            [
                source.title, source.document_type, source.language,
                str(source.source_url or source.local_path or ""), source.license_name,
                source.redistribution_permitted, source.permission_basis,
            ]
        )
    source_sheet.append([])
    source_sheet.append(["Familienstatus kommerzielle Nutzung", questionnaire.is_commercial])
    source_sheet.append(["Version citation", version.source_citation or ""])
    source_sheet.append(["DOI", version.source_doi or ""])
    source_sheet.append(["Erhebungsjahr", version.publication_year])
    source_sheet.append(["Konstrukte", "; ".join(questionnaire.construct_ontology)])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else "A1"
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="28685D")
        for column_cells in sheet.columns:
            width = min(
                60,
                max(
                    12,
                    max(
                        len((str(cell.value or "").splitlines() or [""])[0])
                        for cell in column_cells
                    )
                    + 2,
                ),
            )
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_questionnaire(
    questionnaire: QuestionnaireParent,
    version_id: str,
    export_format: Literal[
        "psymetriq_json", "fhir_json", "item_csv", "redcap_csv", "xlsx", "r_syntax"
    ],
) -> tuple[str, str | bytes]:
    """Export a concrete version and return its recommended extension and content."""
    version = next(
        (candidate for candidate in questionnaire.versions if candidate.version_id == version_id),
        None,
    )
    if version is None:
        raise DataExchangeError(f"Unknown questionnaire version {version_id!r}")
    if export_format == "psymetriq_json":
        return ".json", export_psymetriq_json([questionnaire])
    if export_format == "fhir_json":
        return ".fhir.json", json.dumps(
            export_fhir_questionnaire(questionnaire, version), ensure_ascii=False, indent=2
        ) + "\n"
    if export_format == "item_csv":
        return ".csv", export_item_csv(questionnaire, version)
    if export_format == "redcap_csv":
        return ".csv", export_redcap_data_dictionary(questionnaire, version)
    if export_format == "xlsx":
        return ".xlsx", export_questionnaire_xlsx(questionnaire, version)
    if export_format == "r_syntax":
        return ".R", export_r_syntax(questionnaire, version)
    raise DataExchangeError(f"Unsupported export format: {export_format}")


def select_questionnaire_items(
    questionnaire: QuestionnaireParent,
    version_id: str,
    item_ids: list[str],
    adaptations: dict[str, tuple[str, str]] | None = None,
) -> QuestionnaireParent:
    """Create a validated version subset for a saved questionnaire assembly.

    An empty ``item_ids`` list means the complete version. Scoring algorithms are
    retained only when every target item remains in the selected subset.
    """
    version = next(
        (candidate for candidate in questionnaire.versions if candidate.version_id == version_id),
        None,
    )
    if version is None:
        raise DataExchangeError(f"Unknown questionnaire version {version_id!r}")
    known_ids = {item.item_id for item in version.items}
    requested_ids = set(item_ids) if item_ids else known_ids
    unknown_ids = requested_ids - known_ids
    if unknown_ids:
        raise DataExchangeError(
            f"Saved assembly references unknown item IDs: {', '.join(sorted(unknown_ids))}"
        )
    adaptations = adaptations or {}
    if set(adaptations) - requested_ids:
        raise DataExchangeError("Adaptations must refer to selected items")
    selected_items = []
    adapted_ids = set(adaptations)
    for item in version.items:
        if item.item_id not in requested_ids:
            continue
        if item.item_id in adaptations:
            adapted_text, reason = adaptations[item.item_id]
            previous_notes = item.metadata.notes or ""
            adaptation_note = (
                f"Study-specific adaptation; original catalog item {item.item_id!r} "
                "remains unchanged. This adapted wording has not been psychometrically "
                f"validated. Reason: {reason}"
            )
            item = item.model_copy(
                update={
                    "prompt_text": adapted_text,
                    "metadata": item.metadata.model_copy(
                        update={
                            "notes": "\n".join(
                                value for value in (previous_notes, adaptation_note) if value
                            )
                        }
                    ),
                }
            )
        selected_items.append(item)
    if not selected_items:
        raise DataExchangeError("At least one item must remain selected for export")
    response_set_names = {
        item.response_set_ref for item in selected_items if item.response_set_ref is not None
    }
    selected_scoring = [
        algorithm
        for algorithm in version.scoring_algorithms
        if set(algorithm.target_items).issubset(requested_ids)
        and not (set(algorithm.target_items) & adapted_ids)
    ]
    selected_version = version.model_copy(
        update={
            "items": selected_items,
            "response_sets": {
                name: options
                for name, options in version.response_sets.items()
                if name in response_set_names
            },
            "scoring_algorithms": selected_scoring,
        }
    )
    try:
        return QuestionnaireParent.model_validate(
            questionnaire.model_copy(update={"versions": [selected_version]}).model_dump(
                mode="json"
            )
        )
    except ValidationError as error:
        raise DataExchangeError(
            f"Selected items do not form a valid questionnaire export ({type(error).__name__})"
        ) from None
